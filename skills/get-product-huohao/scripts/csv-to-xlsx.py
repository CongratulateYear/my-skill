import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


HEADERS = {
    "itemId": "商品ID",
    "huohao": "货号",
    "xinghao": "型号",
    "status": "状态",
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="outputs/taobao_properties.csv")
    parser.add_argument("--output", default="outputs/taobao_properties.xlsx")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    with input_path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品属性"

    keys = ["itemId", "huohao", "xinghao", "status"]
    sheet.append([HEADERS[key] for key in keys])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        sheet.append([row.get(key, "") for key in keys])

    for column_index, key in enumerate(keys, start=1):
        values = [HEADERS[key], *[row.get(key, "") for row in rows]]
        width = min(max(len(str(value)) for value in values) + 3, 60)
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    sheet.freeze_panes = "A2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
