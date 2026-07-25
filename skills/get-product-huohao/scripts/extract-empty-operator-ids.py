import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def normalize_header(value):
    return str(value or "").strip().lower()


def is_blank(value):
    return value is None or str(value).strip() == ""


def autosize(sheet):
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(max(len(str(cell.value or "")) for cell in column) + 3, 70)
        sheet.column_dimensions[letter].width = width


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Source Excel workbook with ID and 运营 columns")
    parser.add_argument("--ids-output", default="work/item-ids.txt")
    parser.add_argument("--pending-output", default="outputs/pending_empty_operator_ids.xlsx")
    args = parser.parse_args()

    source = Path(args.input)
    ids_output = Path(args.ids_output)
    pending_output = Path(args.pending_output)

    wb = load_workbook(source, data_only=True)
    ws = wb.active
    headers = [normalize_header(cell.value) for cell in ws[1]]

    id_col = headers.index("id") + 1
    operator_col = headers.index("\u8fd0\u8425") + 1
    title_col = headers.index("\u5546\u54c1\u6807\u9898") + 1 if "\u5546\u54c1\u6807\u9898" in headers else None

    rows = []
    for row_index in range(2, ws.max_row + 1):
        item_id = ws.cell(row_index, id_col).value
        operator = ws.cell(row_index, operator_col).value
        title = ws.cell(row_index, title_col).value if title_col else ""
        if item_id and is_blank(operator):
            rows.append({
                "row": row_index,
                "itemId": str(item_id).strip(),
                "operator": "",
                "title": str(title or "").strip(),
            })

    ids_output.parent.mkdir(parents=True, exist_ok=True)
    ids_output.write_text("\n".join(row["itemId"] for row in rows) + "\n", encoding="utf-8")

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "\u5f85\u67e5\u8be2"
    out_ws.append(["原行号", "商品ID", "运营", "商品标题"])
    for cell in out_ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        out_ws.append([row["row"], row["itemId"], row["operator"], row["title"]])
    out_ws.freeze_panes = "A2"
    autosize(out_ws)

    pending_output.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(pending_output)
    print(f"total_rows={ws.max_row - 1}")
    print(f"pending_ids={len(rows)}")
    print(f"ids_output={ids_output}")
    print(f"pending_output={pending_output}")


if __name__ == "__main__":
    main()
