from __future__ import annotations

import argparse
import re
import shutil
import tempfile
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter


DATE_RE = re.compile(r"(\d+)月(\d+)日")
PDD_SHEET_KEY = "多多买菜"


def date_key(path: Path) -> tuple[int, int, str]:
    match = DATE_RE.search(path.name)
    if match:
        return int(match.group(1)), int(match.group(2)), path.name
    return 99, 99, path.name


def date_label(path: Path) -> str:
    match = DATE_RE.search(path.name)
    if match:
        return f"{int(match.group(1))}月{int(match.group(2))}日"
    return path.stem


def number(value) -> float:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return value
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0


def clone_cell(src, dst) -> None:
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
    if src.comment:
        dst.comment = copy(src.comment)


def setup_sheet_from_template(dst_ws, src_ws, max_row: int, max_col: int) -> None:
    for row in range(1, max_row + 1):
        dst_ws.row_dimensions[row].height = src_ws.row_dimensions[row].height
        for col in range(1, max_col + 1):
            clone_cell(src_ws.cell(row, col), dst_ws.cell(row, col))
            dst_ws.cell(row, col).value = src_ws.cell(row, col).value
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        dst_ws.column_dimensions[letter].width = src_ws.column_dimensions[letter].width
    for merged_range in src_ws.merged_cells.ranges:
        if (
            merged_range.min_row <= max_row
            and merged_range.max_row <= max_row
            and merged_range.min_col <= max_col
            and merged_range.max_col <= max_col
        ):
            dst_ws.merge_cells(str(merged_range))


def find_city_row(ws, city: str) -> int | None:
    city_warehouse = f"{city}仓"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        value = row[0].value
        if isinstance(value, str) and (city in value or city_warehouse in value):
            return row[0].row
    return None


def write_daily_summary_sheet(
    wb_out,
    files: list[Path],
    city: str,
    sheet_name: str,
    max_col: int,
    value_cols: range,
    template_row: int,
    total_template_row: int,
) -> None:
    first = load_workbook(files[0], data_only=True)
    template = first[sheet_name]
    ws = wb_out.create_sheet(sheet_name)

    setup_sheet_from_template(ws, template, 2, max_col)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 16)

    out_row = 3
    daily_rows = []
    for path in files:
        wb = load_workbook(path, data_only=True)
        src = wb[sheet_name]
        city_row = find_city_row(src, city)
        if city_row is None:
            continue

        for col in range(1, max_col + 1):
            clone_cell(template.cell(template_row, col), ws.cell(out_row, col))
        ws.cell(out_row, 1).value = f"{city} {date_label(path)}"
        for col in value_cols:
            ws.cell(out_row, col).value = src.cell(city_row, col).value or 0
        for col in range(max(value_cols) + 1, max_col + 1):
            ws.cell(out_row, col).value = None
        daily_rows.append(out_row)
        out_row += 1

    for col in range(1, max_col + 1):
        clone_cell(template.cell(total_template_row, col), ws.cell(out_row, col))
    ws.cell(out_row, 1).value = "总计"
    for col in value_cols:
        ws.cell(out_row, col).value = sum(number(ws.cell(row, col).value) for row in daily_rows)
    for col in range(max(value_cols) + 1, max_col + 1):
        ws.cell(out_row, col).value = None

    ws.freeze_panes = "A3"


def header_map(ws) -> dict[str, int]:
    headers = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value and value not in headers:
            headers[str(value)] = col
    return headers


def find_pdd_sheet(wb) -> str:
    for name in wb.sheetnames:
        if PDD_SHEET_KEY in name:
            return name
    raise ValueError("未找到多多买菜 sheet")


def city_block_rows(ws, city: str) -> list[dict]:
    headers = header_map(ws)
    item_col = headers.get("平台品类", 4)
    qty_col = headers.get("销售数量", item_col + 1)
    id_col = headers.get("ID", item_col - 1)
    data_cols = {
        "qty": qty_col,
        "sales_price": headers.get("销售单价", qty_col + 1),
        "sales_total": headers.get("销售总额", qty_col + 2),
        "deduction": headers.get("平台扣点", qty_col + 3),
        "settlement_price": headers.get("销售结算单价", qty_col + 4),
        "settlement_total": headers.get("销售结算总额", qty_col + 5),
        "cost_price": headers.get("成本单价", qty_col + 6),
        "cost_total": headers.get("成本总额", qty_col + 7),
    }

    start = None
    for row in range(2, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        if any(value == city or value == f"{city}仓" for value in values):
            start = row
            break
    if start is None:
        return []

    rows = []
    for row in range(start, ws.max_row + 1):
        item = ws.cell(row, item_col).value
        if item in (None, "") and item_col > 1:
            item = ws.cell(row, item_col - 1).value
        if item in (None, ""):
            continue
        rows.append(
            {
                "row": row,
                "style_cols": [
                    1,
                    2,
                    id_col,
                    item_col,
                    data_cols["qty"],
                    data_cols["sales_price"],
                    data_cols["sales_total"],
                    data_cols["deduction"],
                    data_cols["settlement_price"],
                    data_cols["settlement_total"],
                    data_cols["cost_price"],
                    data_cols["cost_total"],
                    data_cols["cost_total"] + 1,
                ],
                "values": [
                    "多多",
                    city,
                    ws.cell(row, id_col).value,
                    item,
                    ws.cell(row, data_cols["qty"]).value,
                    ws.cell(row, data_cols["sales_price"]).value,
                    ws.cell(row, data_cols["sales_total"]).value,
                    ws.cell(row, data_cols["deduction"]).value,
                    ws.cell(row, data_cols["settlement_price"]).value,
                    ws.cell(row, data_cols["settlement_total"]).value,
                    ws.cell(row, data_cols["cost_price"]).value,
                    ws.cell(row, data_cols["cost_total"]).value,
                    "",
                ],
            }
        )
        if str(item).strip() == "合计":
            break
    return rows


def set_border(cell, *, left=None, right=None, top=None, bottom=None) -> None:
    border = cell.border
    cell.border = Border(
        left=left if left is not None else copy(border.left),
        right=right if right is not None else copy(border.right),
        top=top if top is not None else copy(border.top),
        bottom=bottom if bottom is not None else copy(border.bottom),
    )


def apply_pdd_block_borders(ws, start_row: int, end_row: int, max_col: int) -> None:
    thin = Side(style="thin", color="000000")
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            set_border(cell, left=thin, right=thin, top=thin, bottom=thin)
            if col in (1, 2):
                cell.alignment = Alignment(horizontal="center", vertical="center")


def write_pdd_sheet(wb_out, files: list[Path], city: str) -> None:
    first = load_workbook(files[0], data_only=True)
    sheet_name = find_pdd_sheet(first)
    ws = wb_out.create_sheet(sheet_name)
    headers = [
        "平台",
        "基地",
        "ID",
        "平台品类",
        "销售数量",
        "销售单价",
        "销售总额",
        "平台扣点",
        "销售结算单价",
        "销售结算总额",
        "成本单价",
        "成本总额",
        "",
    ]

    first_src = first[sheet_name]
    first_rows = city_block_rows(first_src, city)
    width_cols = [1, 2, *first_rows[0]["style_cols"][2:]] if first_rows else list(range(1, len(headers) + 1))
    for idx, src_col in enumerate(width_cols, 1):
        ws.column_dimensions[get_column_letter(idx)].width = first_src.column_dimensions[get_column_letter(src_col)].width

    row = 1
    for path in files:
        wb = load_workbook(path, data_only=True)
        src = wb[find_pdd_sheet(wb)]
        rows = city_block_rows(src, city)
        if not rows:
            continue

        ws.row_dimensions[row].height = src.row_dimensions[1].height
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        clone_cell(src.cell(1, 1), ws.cell(row, 1))
        ws.cell(row, 1).value = date_label(path)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        row += 1

        ws.row_dimensions[row].height = src.row_dimensions[1].height
        header_cols = [1, 2, *rows[0]["style_cols"][2:]]
        for col, (header, src_col) in enumerate(zip(headers, header_cols), 1):
            if src_col <= src.max_column:
                clone_cell(src.cell(1, src_col), ws.cell(row, col))
            ws.cell(row, col).value = header or None
        row += 1

        block_start = row
        for record in rows:
            ws.row_dimensions[row].height = src.row_dimensions[record["row"]].height
            for col, value in enumerate(record["values"], 1):
                cell = ws.cell(row, col, value)
                src_col = record["style_cols"][col - 1]
                if src_col <= src.max_column:
                    clone_cell(src.cell(record["row"], src_col), cell)
                    cell.value = value
                if col in (1, 2) and row > block_start:
                    cell.value = None
            row += 1

        block_end = row - 1
        clone_cell(src.cell(2, 1), ws.cell(block_start, 1))
        ws.cell(block_start, 1).value = "多多"
        ws.cell(block_start, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(block_start, 2).alignment = Alignment(horizontal="center", vertical="center")
        apply_pdd_block_borders(ws, block_start, block_end, len(headers))
        ws.merge_cells(start_row=block_start, start_column=1, end_row=block_end, end_column=1)
        ws.merge_cells(start_row=block_start, start_column=2, end_row=block_end, end_column=2)

        blank_src_row = min(rows[-1]["row"] + 1, src.max_row)
        ws.row_dimensions[row].height = src.row_dimensions[blank_src_row].height
        for col, src_col in enumerate(header_cols, 1):
            if src_col <= src.max_column:
                clone_cell(src.cell(blank_src_row, src_col), ws.cell(row, col))
            ws.cell(row, col).value = None
        row += 1

    ws.freeze_panes = "A3"


def collect_inputs(paths: list[Path], temp_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in paths:
        if path.is_dir():
            files.extend(path.rglob("*.xlsx"))
        elif path.suffix.lower() == ".zip":
            extract_dir = temp_dir / path.stem
            extract_dir.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(path) as zf:
                zf.extractall(extract_dir)
            files.extend(extract_dir.rglob("*.xlsx"))
        elif path.suffix.lower() == ".xlsx":
            files.append(path)
    clean = [
        path
        for path in files
        if not path.name.startswith("~$")
        and "周报汇总表" not in path.name
        and path.is_file()
    ]
    return sorted(clean, key=date_key)


def build_report(input_paths: list[Path], city: str, output_dir: Path) -> Path:
    temp_dir = Path(tempfile.mkdtemp(prefix="weekly-report-summary-"))
    try:
        files = collect_inputs(input_paths, temp_dir)
        if not files:
            raise SystemExit("未找到可处理的 .xlsx 文件")

        wb = Workbook()
        write_daily_summary_sheet(wb, files, city, "营收概况", 8, range(2, 7), 4, 8)
        write_daily_summary_sheet(wb, files, city, "基地汇总", 17, range(2, 18), 4, 8)
        write_pdd_sheet(wb, files, city)
        for ws in list(wb.worksheets):
            if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
                wb.remove(ws)

        output_dir.mkdir(parents=True, exist_ok=True)
        output = output_dir / f"{city} 周报汇总表.xlsx"
        wb.save(output)
        return output
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build city weekly summary workbook.")
    parser.add_argument("inputs", nargs="+", type=Path, help="xlsx files, directories, or zip archives")
    parser.add_argument("--city", required=True, help="city name, e.g. 苏州")
    parser.add_argument("--output-dir", type=Path, default=Path.cwd() / "outputs" / "weekly-report-summary")
    args = parser.parse_args()

    output = build_report(args.inputs, args.city.strip(), args.output_dir)
    print(output)


if __name__ == "__main__":
    main()
