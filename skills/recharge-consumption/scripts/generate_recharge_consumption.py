#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate the fixed-format recharge/consumption workbook from Sheet1."""

from __future__ import annotations

import argparse
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook


OUTPUT_NAME = "\u5145\u503c\u6d88\u8017.xlsx"
SHEET_NAME = "\u5145\u503c\u6d88\u8017"

DATE_COL = "\u4e1a\u52a1\u53d1\u751f\u65e5\u671f"
SHOP_COL = "\u5e97\u94fa"
TRADE_COL = "\u4ea4\u6613\u7c7b\u578b"
AMOUNT_COL = "\u91d1\u989d"
RATIO_COL = "\u62b5\u7528\u5238\u5360\u6bd4"
TOTAL_LABEL = "\u5408\u8ba1"

HEADERS = [
    DATE_COL,
    SHOP_COL,
    "\u3010\u534f\u8bae\u6263\u6b3e\u3011- \u81ea\u52a8\u5145\u503c",
    "\u5145\u503c\u52a0\u7801\u6d88\u8017",
    "\u6d41\u91cf\u5238\u62b5\u6263",
    "\u6295\u6d41\u63a8\u5e7f\u5145\u503c",
    "\u624b\u52a8\u5145\u503c",
    "\u9884\u7b97\u6d88\u8d39",
    RATIO_COL,
]
AMOUNT_HEADERS = HEADERS[2:-1]
REQUIRED_SHEET1_COLUMNS = [DATE_COL, SHOP_COL, TRADE_COL, AMOUNT_COL]


def decode_escaped(value: str | None) -> str | None:
    if value is None:
        return None
    return value.encode("ascii").decode("unicode_escape")


def norm_date(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()[:10]


def to_decimal(value) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        return Decimal("0")


def resolve_source(args: argparse.Namespace) -> Path:
    if args.input:
        source = Path(args.input).expanduser()
        if source.exists():
            return source.resolve()

        wildcard_name = source.name.replace("?", "*")
        if wildcard_name != source.name and source.parent.exists():
            matches = sorted(source.parent.glob(wildcard_name))
            if len(matches) == 1:
                return matches[0].resolve()
            if matches:
                raise FileNotFoundError(
                    "Input path was corrupted and matched multiple files: "
                    + "; ".join(str(path) for path in matches)
                )

        raise FileNotFoundError(f"Source workbook does not exist: {source}")

    if not args.source_dir:
        raise ValueError("Provide --input or --source-dir")

    source_dir = Path(args.source_dir).expanduser()
    if not source_dir.exists():
        raise FileNotFoundError(f"Source directory does not exist: {source_dir}")

    candidates = [path for path in source_dir.glob(args.glob) if path.is_file()]
    name_contains = args.name_contains or decode_escaped(args.name_contains_escaped)
    if name_contains:
        candidates = [path for path in candidates if name_contains in path.name]

    if len(candidates) != 1:
        raise FileNotFoundError(
            f"Expected exactly one source workbook, found {len(candidates)}: "
            + "; ".join(str(path) for path in sorted(candidates))
        )
    return candidates[0].resolve()


def make_workbook(source: Path, output_dir: Path) -> Path:
    source_wb = load_workbook(source, data_only=False)
    if "Sheet1" not in source_wb.sheetnames:
        raise ValueError("Source workbook must contain Sheet1")

    ws1 = source_wb["Sheet1"]
    sheet1_headers = [ws1.cell(1, col).value for col in range(1, ws1.max_column + 1)]
    sheet1_index = {header: col + 1 for col, header in enumerate(sheet1_headers)}
    missing = [name for name in REQUIRED_SHEET1_COLUMNS if name not in sheet1_index]
    if missing:
        raise ValueError("Sheet1 is missing required columns: " + ", ".join(missing))

    amount_by_key = defaultdict(Decimal)
    date_shops: dict[str, set[str]] = {}

    for row_num in range(2, ws1.max_row + 1):
        date = norm_date(ws1.cell(row_num, sheet1_index[DATE_COL]).value)
        shop_value = ws1.cell(row_num, sheet1_index[SHOP_COL]).value
        trade_value = ws1.cell(row_num, sheet1_index[TRADE_COL]).value
        if not date or shop_value is None or trade_value is None:
            continue

        shop = str(shop_value).strip()
        trade = str(trade_value).strip()
        if not shop or not trade:
            continue

        date_shops.setdefault(date, set()).add(shop)
        if trade not in AMOUNT_HEADERS:
            continue

        amount = to_decimal(ws1.cell(row_num, sheet1_index[AMOUNT_COL]).value)
        amount_by_key[(date, shop, trade)] += amount

    output_wb = Workbook()
    ws = output_wb.active
    ws.title = SHEET_NAME

    for col, header in enumerate(HEADERS, start=1):
        ws.cell(1, col, header)

    row_num = 2
    dates = sorted(date_shops)
    for date_index, date in enumerate(dates):
        totals = {header: Decimal("0") for header in AMOUNT_HEADERS}

        for shop in sorted(date_shops[date]):
            ws.cell(row_num, 1, date)
            ws.cell(row_num, 2, shop)
            for col, header in enumerate(HEADERS[2:], start=3):
                if header == RATIO_COL:
                    ws.cell(row_num, col, None)
                    continue

                value = amount_by_key.get((date, shop, header), Decimal("0"))
                totals[header] += value
                ws.cell(row_num, col, float(value))
            row_num += 1

        ws.cell(row_num, 1, date)
        ws.cell(row_num, 2, TOTAL_LABEL)
        for col, header in enumerate(HEADERS[2:], start=3):
            ws.cell(row_num, col, None if header == RATIO_COL else float(totals[header]))
        row_num += 1

        if date_index != len(dates) - 1:
            row_num += 2

    for row in ws.iter_rows():
        for cell in row:
            cell.number_format = "General"

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_NAME
    output_wb.save(output_path)
    return output_path


def validate_workbook(output_path: Path) -> dict[str, int]:
    if not output_path.exists():
        raise FileNotFoundError(f"Output workbook does not exist: {output_path}")

    workbook = load_workbook(output_path, data_only=False)
    errors = []
    if workbook.sheetnames != [SHEET_NAME]:
        errors.append(f"expected one sheet named {SHEET_NAME!r}, found {workbook.sheetnames!r}")

    worksheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
    headers = [worksheet.cell(1, column).value for column in range(1, len(HEADERS) + 1)]
    if headers != HEADERS:
        errors.append("fixed headers do not match")

    number_formats = sorted({cell.number_format for row in worksheet.iter_rows() for cell in row})
    if number_formats != ["General"]:
        errors.append(f"expected only General number format, found {number_formats!r}")

    nonblank_rows = [
        row
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]
    total_rows = [row for row in nonblank_rows if len(row) > 1 and row[1] == TOTAL_LABEL]
    if nonblank_rows and not total_rows:
        errors.append("no total rows found")

    if errors:
        raise ValueError("Validation failed: " + "; ".join(errors))

    return {
        "max_row": worksheet.max_row,
        "max_column": worksheet.max_column,
        "nonblank_rows": len(nonblank_rows),
        "total_rows": len(total_rows),
    }


def print_validation_summary(summary: dict[str, int]) -> None:
    print(
        "validation: ok; "
        f"rows={summary['max_row']}; "
        f"columns={summary['max_column']}; "
        f"nonblank_rows={summary['nonblank_rows']}; "
        f"total_rows={summary['total_rows']}"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate recharge/consumption summary workbook from Sheet1 data."
    )
    parser.add_argument("--input", help="Source Excel workbook path.")
    parser.add_argument(
        "--output-dir",
        help="Output directory. Defaults to the source workbook directory.",
    )
    parser.add_argument(
        "--source-dir",
        help="Directory to search when Chinese source paths cannot be passed safely.",
    )
    parser.add_argument(
        "--glob",
        default="*.xlsx",
        help="Workbook glob used with --source-dir. Defaults to *.xlsx.",
    )
    parser.add_argument(
        "--name-contains",
        help="Filter source candidates by literal text in the filename.",
    )
    parser.add_argument(
        "--name-contains-escaped",
        help=r"ASCII-safe filename filter using Python Unicode escapes, e.g. \u5168\u90e8\u516c\u53f8.",
    )
    parser.add_argument(
        "--skip-validation",
        action="store_true",
        help="Skip the built-in output workbook validation.",
    )
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help="Validate the existing output workbook in --output-dir and exit.",
    )
    args = parser.parse_args()

    if args.validate_only:
        if not args.output_dir:
            raise ValueError("Provide --output-dir with --validate-only")
        output_path = Path(args.output_dir).expanduser().resolve() / OUTPUT_NAME
        print_validation_summary(validate_workbook(output_path))
        return

    source = resolve_source(args)
    output_dir = Path(args.output_dir).expanduser().resolve() if args.output_dir else source.parent
    output_path = make_workbook(source, output_dir)
    print(output_path)
    if not args.skip_validation:
        print_validation_summary(validate_workbook(output_path))


if __name__ == "__main__":
    main()
